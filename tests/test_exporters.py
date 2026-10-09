"""CSV/XLSX output: file names, layout, formatting and masking."""

import csv
import io

import pytest
from openpyxl import load_workbook

from data_health_check import exporters, profiler


def column(name, data_type="string", label=None, samples=("a", "b"), **stats):
    base = {
        "column_name": name,
        "data_type": data_type,
        "data_label": label,
        "data_label_confidence": 0.99 if label else None,
        "null_count": 0,
        "null_percent": 0.0,
        "unique_count": 2,
        "unique_ratio": 1.0,
        "min": None,
        "max": None,
        "mean": None,
        "median": None,
        "stddev": None,
        "sample_values": list(samples),
        "rows_profiled": 2,
    }
    base.update(stats)
    return base


def make_profile(file_format="csv", mode="full", sheets=None):
    sheets = sheets or [
        profiler.SheetProfile(
            sheet_name=None if file_format == "csv" else "People",
            row_count=2,
            column_count=3,
            rows_profiled=2,
            columns=[
                column("email", label="EMAIL_ADDRESS", samples=["a@example.com"]),
                column("ssn", label="UNKNOWN|SSN", samples=["900-00-0001"]),
                column("age", "int", "INTEGER", ["30", "41"], min=30, max=41),
            ],
        )
    ]
    return profiler.DatasetProfile(
        file_name=f"people.{file_format}",
        file_format=file_format,
        mode=mode,
        profiled_at="2026-01-02T03:04:05Z",
        sheets=sheets,
    )


def read_csv(data: bytes) -> list[dict]:
    return list(csv.DictReader(io.StringIO(data.decode("utf-8-sig"))))


@pytest.mark.parametrize(
    ("name", "expected"),
    [
        ("customers.csv", "customers_profile.csv"),
        ("Q3 Report.XLSX", "Q3 Report_profile.xlsx"),
        ("my.data.v2.csv", "my.data.v2_profile.csv"),
    ],
)
def test_output_file_name(name, expected):
    assert exporters.output_file_name(name) == expected


# --- CSV ---------------------------------------------------------------------


def test_csv_in_gives_csv_out():
    exported = exporters.export_profile(make_profile("csv"))
    assert exported.file_name == "people_profile.csv"
    assert exported.mime_type == "text/csv"
    assert exported.data.startswith(b"\xef\xbb\xbf")  # BOM so Excel reads UTF-8


def test_csv_has_one_row_per_column_and_all_fields():
    rows = read_csv(exporters.to_csv_bytes(make_profile("csv")))
    assert [row["column_name"] for row in rows] == ["email", "ssn", "age"]
    assert list(rows[0]) == list(profiler.COLUMN_FIELDS)
    assert rows[2]["min"] == "30"
    assert rows[0]["min"] == ""  # not applicable -> empty


def test_csv_masks_sensitive_samples_by_default():
    rows = read_csv(exporters.to_csv_bytes(make_profile("csv")))
    assert rows[0]["sample_values"] == "[masked: EMAIL_ADDRESS]"
    assert rows[1]["sample_values"] == "[masked: SSN]"
    assert rows[2]["sample_values"] == "30; 41"


def test_csv_can_reveal_samples():
    data = exporters.to_csv_bytes(make_profile("csv"), reveal_sensitive=True)
    assert read_csv(data)[0]["sample_values"] == "a@example.com"


def test_lite_csv_has_no_label_columns():
    rows = read_csv(exporters.to_csv_bytes(make_profile("csv", mode="lite")))
    assert "data_label" not in rows[0]
    assert "data_label_confidence" not in rows[0]


def test_csv_neutralizes_formulas_but_not_negative_numbers():
    sheet = profiler.SheetProfile(
        None,
        3,
        2,
        3,
        [
            column("=cmd", samples=['=HYPERLINK("http://x")', "@SUM(A1)"]),
            column("balance", "float", "FLOAT", ["-12.5", "3"], min=-12.5),
        ],
    )
    rows = read_csv(exporters.to_csv_bytes(make_profile("csv", sheets=[sheet])))
    assert rows[0]["column_name"] == "'=cmd"
    assert rows[0]["sample_values"].startswith("'=HYPERLINK")
    assert rows[1]["sample_values"] == "-12.5; 3"
    assert rows[1]["min"] == "-12.5"


def test_real_csv_profile_round_trips(csv_profile_lite):
    rows = read_csv(exporters.export_profile(csv_profile_lite).data)
    assert len(rows) == 16
    balance = next(row for row in rows if row["column_name"] == "account_balance")
    assert balance["null_count"] == "16"


# --- XLSX --------------------------------------------------------------------


def load(data: bytes):
    return load_workbook(io.BytesIO(data))


def test_xlsx_in_gives_xlsx_out_with_expected_sheets():
    exported = exporters.export_profile(make_profile("xlsx"))
    assert exported.file_name == "people_profile.xlsx"
    assert exported.mime_type.endswith("spreadsheetml.sheet")
    assert load(exported.data).sheetnames == ["Summary", "People", "Field guide"]


def test_xlsx_summary_sheet():
    sheet = load(exporters.to_xlsx_bytes(make_profile("xlsx")))["Summary"]
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0] == exporters.SUMMARY_HEADERS
    assert rows[1][:7] == (
        "people.xlsx",
        "People",
        2,
        3,
        2,
        "2026-01-02T03:04:05Z",
        "full",
    )


def test_xlsx_headers_are_bold_and_frozen_with_sensible_widths():
    workbook = load(exporters.to_xlsx_bytes(make_profile("xlsx")))
    for sheet in workbook.worksheets:
        assert sheet.freeze_panes == "A2"
        assert all(cell.font.bold for cell in sheet[1])
        widths = [dim.width for dim in sheet.column_dimensions.values()]
        assert widths and all(8 <= width <= 100 for width in widths)


def test_xlsx_profile_sheet_masks_by_default():
    sheet = load(exporters.to_xlsx_bytes(make_profile("xlsx")))["People"]
    rows = list(sheet.iter_rows(values_only=True))
    header = rows[0]
    samples = [row[header.index("sample_values")] for row in rows[1:]]
    assert samples == ["[masked: EMAIL_ADDRESS]", "[masked: SSN]", "30; 41"]
    assert rows[3][header.index("min")] == 30


def test_xlsx_field_guide_explains_every_field():
    sheet = load(exporters.to_xlsx_bytes(make_profile("xlsx")))["Field guide"]
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0] == ("Field", "What it means")
    assert [row[0] for row in rows[1:]] == list(profiler.COLUMN_FIELDS)
    assert all(row[1] for row in rows[1:])


def test_xlsx_lite_field_guide_matches_lite_fields():
    workbook = load(exporters.to_xlsx_bytes(make_profile("xlsx", mode="lite")))
    fields = [row[0] for row in workbook["Field guide"].iter_rows(values_only=True)]
    assert "data_label" not in fields
    assert "data_label" not in next(workbook["People"].iter_rows(values_only=True))


def test_xlsx_text_that_looks_like_a_formula_stays_text():
    sheet = profiler.SheetProfile(
        "S", 1, 1, 1, [column("=1+1", samples=["=SUM(A1:A2)", "\x00\x07bell"])]
    )
    data = exporters.to_xlsx_bytes(make_profile("xlsx", sheets=[sheet]))
    cells = load(data)["S"]
    assert cells["A2"].value == "=1+1"
    assert cells["A2"].data_type == "s"
    assert "\x07" not in cells["N2"].value


def test_xlsx_sheet_names_are_valid_and_unique():
    sheets = [
        profiler.SheetProfile(name, 1, 1, 1, [column("x")])
        for name in ["Summary", "summary", "Field guide", "a/b:c*d?" + "x" * 40]
    ]
    names = load(
        exporters.to_xlsx_bytes(make_profile("xlsx", sheets=sheets))
    ).sheetnames
    assert names[0] == "Summary" and names[-1] == "Field guide"
    profile_names = names[1:-1]
    assert len({name.lower() for name in names}) == len(names)
    assert all(len(name) <= 31 for name in names)
    assert not any(ch in name for name in profile_names for ch in "[]:*?/\\")


def test_real_xlsx_profile(xlsx_profile_lite):
    workbook = load(exporters.export_profile(xlsx_profile_lite).data)
    assert workbook.sheetnames == ["Summary", "Customers", "Orders", "Field guide"]
    assert workbook["Customers"].max_row == 17  # header + 16 columns
    assert workbook["Orders"].max_row == 10


@pytest.mark.full_mode
def test_full_mode_xlsx_export_masks_sensitive_columns(xlsx_profile_full):
    workbook = load(exporters.export_profile(xlsx_profile_full).data)
    rows = list(workbook["Customers"].iter_rows(values_only=True))
    header = rows[0]
    by_name = {row[0]: row for row in rows[1:]}
    assert by_name["ssn"][header.index("data_label")] == "SSN"
    assert by_name["ssn"][header.index("sample_values")] == "[masked: SSN]"
    assert "[masked" not in by_name["plan"][header.index("sample_values")]
