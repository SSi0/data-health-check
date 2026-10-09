"""File type, size, encoding and damaged-file handling (no TensorFlow needed)."""

import codecs
import io

import pytest
from openpyxl import Workbook

from conftest import COMPANY_XLSX, CUSTOMERS_CSV, PASSWORD_XLSX
from data_health_check import profiler


def profile_lite(data: bytes, name: str) -> profiler.DatasetProfile:
    return profiler.profile_file(data, name, detect_pii=False)


def workbook_bytes(build) -> bytes:
    workbook = Workbook()
    build(workbook)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


@pytest.mark.parametrize(
    ("name", "expected"),
    [("a.csv", "csv"), ("A.CSV", "csv"), ("b.xlsx", "xlsx"), ("dir/c.XLSX", "xlsx")],
)
def test_detect_format(name, expected):
    assert profiler.detect_format(name) == expected


@pytest.mark.parametrize("name", ["data.txt", "data.xls", "data.json", "data"])
def test_unsupported_file_types_are_rejected(name):
    with pytest.raises(profiler.UnsupportedFileTypeError, match=r"\.csv or \.xlsx"):
        profile_lite(b"a,b\n1,2\n", name)


def test_file_over_the_size_limit_is_rejected():
    with pytest.raises(profiler.FileTooLargeError, match="25 MB"):
        profiler.check_size(profiler.MAX_FILE_BYTES + 1)
    profiler.check_size(profiler.MAX_FILE_BYTES)


@pytest.mark.parametrize("data", [b"", b"   \n\n"])
def test_empty_csv(data):
    with pytest.raises(profiler.EmptyFileError, match="empty"):
        profile_lite(data, "empty.csv")


def test_csv_with_only_a_header():
    with pytest.raises(profiler.EmptyFileError, match="no data rows"):
        profile_lite(b"name,email\n", "header.csv")


def test_empty_xlsx_file():
    with pytest.raises(profiler.EmptyFileError):
        profile_lite(b"", "empty.xlsx")


def test_workbook_without_data():
    data = workbook_bytes(lambda wb: wb.create_sheet("Blank"))
    with pytest.raises(profiler.EmptyFileError, match="None of the sheets"):
        profile_lite(data, "blank.xlsx")


def test_password_protected_xlsx():
    with pytest.raises(profiler.PasswordProtectedError, match="password"):
        profile_lite(PASSWORD_XLSX.read_bytes(), "locked.xlsx")


def test_legacy_xls_renamed_to_xlsx():
    fake_xls = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 512
    with pytest.raises(profiler.UnsupportedFileTypeError, match=r"\.xls"):
        profile_lite(fake_xls, "old.xlsx")


def test_damaged_xlsx():
    data = COMPANY_XLSX.read_bytes()[:2000]  # truncated zip
    with pytest.raises(profiler.UnreadableFileError, match="couldn't open"):
        profile_lite(data, "damaged.xlsx")


def test_csv_file_renamed_to_xlsx():
    with pytest.raises(profiler.UnreadableFileError, match="isn't a valid .xlsx"):
        profile_lite(CUSTOMERS_CSV.read_bytes(), "customers.xlsx")


def test_xlsx_file_renamed_to_csv():
    with pytest.raises(profiler.UnreadableFileError, match="Excel file"):
        profile_lite(COMPANY_XLSX.read_bytes(), "company.csv")


def test_binary_file_named_csv():
    with pytest.raises(profiler.UnreadableFileError, match="doesn't look like a CSV"):
        profile_lite(b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR", "image.csv")


def test_unreadable_encoding():
    # 0x81 and 0x8d are invalid in both UTF-8 and Windows-1252.
    with pytest.raises(profiler.EncodingError, match="CSV UTF-8"):
        profile_lite(b"name,city\n\x81\x8d,x\n", "weird.csv")


@pytest.mark.parametrize(
    ("encoded", "label"),
    [
        ("name,city\nJosé,Zürich\nRenée,Montréal\n".encode(), "utf-8"),
        (codecs.BOM_UTF8 + "name,city\nJosé,Zürich\n".encode(), "utf-8 with BOM"),
        ("name,city\nJosé,Zürich\n".encode("cp1252"), "windows-1252"),
        ("name,city\nJosé,Zürich\n".encode("utf-16"), "utf-16"),
    ],
)
def test_common_encodings_are_read(encoded, label):
    profile = profile_lite(encoded, "accents.csv")
    columns = [column["column_name"] for column in profile.sheets[0].columns]
    assert columns == ["name", "city"], label
    assert "José" in profile.sheets[0].columns[0]["sample_values"]


@pytest.mark.parametrize("delimiter", [";", "\t", "|"])
def test_other_delimiters_are_detected(delimiter):
    rows = ["id{0}amount", "1{0}2.5", "2{0}3.5", "3{0}4.5"]
    data = "\n".join(row.format(delimiter) for row in rows).encode()
    profile = profile_lite(data, "semi.csv")
    assert [c["column_name"] for c in profile.sheets[0].columns] == ["id", "amount"]


def test_ragged_csv_gives_a_friendly_error():
    data = b"a,b\n1,2\n3,4,5,6\n"
    with pytest.raises(profiler.UnreadableFileError, match="don't line up"):
        profile_lite(data, "ragged.csv")


def test_blank_excel_rows_and_unnamed_empty_columns_are_ignored():
    def build(wb):
        sheet = wb.active
        sheet.title = "Data"
        sheet.append(["id", "value"])
        sheet.append([1, "a"])
        sheet.append([None, None])
        sheet.append([2, "b"])
        sheet["D10"].value = None
        sheet["D10"].number_format = "0.00"  # formatted but empty cell

    sheet = profile_lite(workbook_bytes(build), "gaps.xlsx").sheets[0]
    assert sheet.row_count == 2
    assert [c["column_name"] for c in sheet.columns] == ["id", "value"]


def test_empty_sheet_next_to_data_is_reported_not_profiled():
    def build(wb):
        wb.active.title = "Data"
        wb.active.append(["id"])
        wb.active.append([1])
        wb.create_sheet("Empty")

    profile = profile_lite(workbook_bytes(build), "two.xlsx")
    data, empty = profile.sheets
    assert (data.sheet_name, data.row_count) == ("Data", 1)
    assert (empty.sheet_name, empty.row_count, empty.columns) == ("Empty", 0, [])
    assert "nothing to profile" in empty.note
