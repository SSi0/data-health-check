"""The MCP server: tool listing, JSON output, masking, paths and errors."""

import asyncio
import csv
import io
import json
import logging
import os
import re
import shutil
import sys
from pathlib import Path

import pytest
from mcp import Client, StdioServerParameters
from mcp.client.stdio import get_default_environment
from openpyxl import load_workbook

import data_health_check
from conftest import (
    COMPANY_XLSX,
    CUSTOMERS_CSV,
    PASSWORD_XLSX,
    digits_only_ids_csv,
    mixed_contacts_csv,
    without_header,
)
from data_health_check import masking, mcp_server, profiler

SAMPLE_FILES = [CUSTOMERS_CSV, COMPANY_XLSX]
# (row_count, column_count, sheet_count) of the sample files.
SUMMARIES = {".csv": (200, 16, 1), ".xlsx": (400, 25, 2)}
posix_only = pytest.mark.skipif(os.name == "nt", reason="needs POSIX links")
# Raw values that must never reach an agent from the sample files.
RAW_SENSITIVE = [
    re.compile(r"@example\.(com|org)"),  # emails
    re.compile(r"555-01\d\d"),  # phone numbers
    re.compile(r"\b9\d\d-\d\d-\d{4}\b"),  # SSNs
    re.compile(r"192\.0\.2\.\d+"),  # IP addresses
    re.compile(r"\d{4} \d{4} \d{4} \d{4}"),  # card numbers
]


def call(tool: str, target=mcp_server.server, **arguments):
    """Call a tool through a real MCP client session and return the result."""

    async def go():
        async with Client(target, read_timeout_seconds=300) as client:
            return await client.call_tool(tool, arguments)

    return asyncio.run(go())


def ok_json(result) -> dict:
    assert not result.is_error, result.content[0].text
    data = json.loads(result.content[0].text)  # the text content is valid JSON
    assert result.structured_content == data
    return data


def error_text(result) -> str:
    assert result.is_error
    return result.content[0].text


def columns(data: dict) -> list[dict]:
    return [column for sheet in data["sheets"] for column in sheet["columns"]]


def assert_no_raw_sensitive_values(text: str) -> None:
    for pattern in RAW_SENSITIVE:
        assert not pattern.search(text), pattern.pattern


# --- Tool listing ---------------------------------------------------------------


def test_tools_are_listed_with_agent_facing_descriptions():
    async def go():
        async with Client(mcp_server.server) as client:
            return (await client.list_tools()).tools

    tools = {tool.name: tool for tool in asyncio.run(go())}
    assert set(tools) == {"profile_dataset", "export_profile"}

    profile = tools["profile_dataset"]
    props = profile.input_schema["properties"]
    assert profile.input_schema["required"] == ["file_path"]
    assert props["detect_pii"]["default"] is True
    assert props["include_samples"]["default"] is False
    assert all(prop.get("description") for prop in props.values())
    assert "Use this" in profile.description
    assert "[masked: <LABEL>]" in profile.description
    assert profile.annotations.read_only_hint is True

    export = tools["export_profile"]
    props = export.input_schema["properties"]
    assert export.input_schema["required"] == ["file_path"]
    assert props["output_path"]["default"] is None
    assert props["detect_pii"]["default"] is True
    assert "[masked: <LABEL>]" in export.description
    assert export.annotations.read_only_hint is False
    assert export.annotations.destructive_hint is True  # replaces earlier profiles

    # The agent is told what is guaranteed, not that nothing can leak.
    for text in (profile.description, mcp_server.INSTRUCTIONS):
        assert "never returned" not in text
        assert "can miss" in text


# --- profile_dataset ------------------------------------------------------------


@pytest.mark.parametrize("path", SAMPLE_FILES, ids=lambda p: p.suffix)
def test_lite_profile_is_valid_json(path):
    data = ok_json(call("profile_dataset", file_path=str(path), detect_pii=False))
    assert data["file_name"] == path.name
    assert data["mode"] == "lite"
    summary = (data["row_count"], data["column_count"], data["sheet_count"])
    assert summary == SUMMARIES[path.suffix]
    for sheet in data["sheets"]:
        assert sheet["column_count"] == len(sheet["columns"])
    for column in columns(data):
        assert {"column_name", "data_type", "null_count", "unique_count"} <= set(column)
        assert "data_label" not in column
        assert "sample_values" not in column  # include_samples defaults to False


@pytest.mark.parametrize("path", SAMPLE_FILES, ids=lambda p: p.suffix)
def test_lite_mode_never_returns_samples(path):
    result = call(
        "profile_dataset", file_path=str(path), detect_pii=False, include_samples=True
    )
    data = ok_json(result)
    assert {c["sample_values"] for c in columns(data)} == {
        masking.LITE_MODE_PLACEHOLDER
    }
    assert_no_raw_sensitive_values(result.content[0].text)


def test_home_directory_paths_are_expanded(tmp_path, monkeypatch):
    shutil.copy(CUSTOMERS_CSV, tmp_path / "customers.csv")
    monkeypatch.setenv("HOME", str(tmp_path))
    monkeypatch.setenv("USERPROFILE", str(tmp_path))  # Windows
    data = ok_json(
        call("profile_dataset", file_path="~/customers.csv", detect_pii=False)
    )
    assert data["row_count"] == 200


@pytest.mark.full_mode
@pytest.mark.parametrize("path", SAMPLE_FILES, ids=lambda p: p.suffix)
def test_full_profile_is_valid_json_with_labels_and_no_samples(path):
    result = call("profile_dataset", file_path=str(path))
    data = ok_json(result)
    assert data["mode"] == "full"
    labels = {c["column_name"]: c.get("data_label") for c in columns(data)}
    assert labels["ssn"] == "SSN"
    assert labels["email"] == "EMAIL_ADDRESS"
    assert all("sample_values" not in c for c in columns(data))
    assert_no_raw_sensitive_values(result.content[0].text)


@pytest.mark.full_mode
@pytest.mark.parametrize("path", SAMPLE_FILES, ids=lambda p: p.suffix)
def test_full_profile_masks_sensitive_samples(path):
    result = call("profile_dataset", file_path=str(path), include_samples=True)
    data = ok_json(result)
    for column in columns(data):
        label = column.get("data_label")
        if masking.is_sensitive(label):
            assert column["sample_values"] == masking.mask_placeholder(label)
            assert column["sample_values"].startswith("[masked: ")
        else:
            assert "[masked" not in column["sample_values"]
    plan = next(c for c in columns(data) if c["column_name"] == "plan")
    assert plan["sample_values"].count(";") == 4  # non-sensitive: 5 raw values
    assert_no_raw_sensitive_values(result.content[0].text)


# --- export_profile -------------------------------------------------------------


@pytest.fixture
def workdir(tmp_path) -> Path:
    for path in SAMPLE_FILES:
        shutil.copy(path, tmp_path / path.name)
    return tmp_path


def exported(result) -> Path:
    assert not result.is_error, result.content[0].text
    path = Path(result.structured_content["result"])
    assert path.is_absolute() and path.is_file()
    return path


def test_export_csv_next_to_the_input(workdir):
    out = exported(
        call(
            "export_profile", file_path=str(workdir / "customers.csv"), detect_pii=False
        )
    )
    assert out == workdir / "customers_profile.csv"
    rows = list(csv.DictReader(io.StringIO(out.read_text(encoding="utf-8-sig"))))
    assert len(rows) == 16
    assert {row["sample_values"] for row in rows} == {masking.LITE_MODE_PLACEHOLDER}


def test_export_xlsx_next_to_the_input(workdir):
    out = exported(
        call(
            "export_profile",
            file_path=str(workdir / "company_data.xlsx"),
            detect_pii=False,
        )
    )
    assert out == workdir / "company_data_profile.xlsx"
    names = load_workbook(out).sheetnames
    assert names == ["Summary", "Customers", "Orders", "Field guide"]


def test_export_to_a_folder_or_relative_path(workdir):
    (workdir / "reports").mkdir()
    source = str(workdir / "customers.csv")
    to_folder = exported(
        call(
            "export_profile",
            file_path=source,
            output_path=str(workdir / "reports"),
            detect_pii=False,
        )
    )
    assert to_folder == workdir / "reports" / "customers_profile.csv"
    relative = exported(
        call(
            "export_profile",
            file_path=source,
            output_path="reports/q3.csv",
            detect_pii=False,
        )
    )
    assert relative == workdir / "reports" / "q3.csv"


def test_export_replaces_an_earlier_profile(workdir):
    source = str(workdir / "customers.csv")
    first = exported(call("export_profile", file_path=source, detect_pii=False))
    first.write_text(first.read_text(encoding="utf-8-sig") + "edited,later\n")
    again = exported(call("export_profile", file_path=source, detect_pii=False))
    assert again == first
    assert "edited,later" not in again.read_text(encoding="utf-8-sig")


def test_export_can_rerun_with_a_custom_name(workdir):
    source = str(workdir / "company_data.xlsx")
    for _ in range(2):
        out = exported(
            call(
                "export_profile",
                file_path=source,
                output_path="q3_report.xlsx",
                detect_pii=False,
            )
        )
    assert out == workdir / "q3_report.xlsx"


@pytest.mark.parametrize("name", ["customers_profile.csv", "company_data_profile.xlsx"])
def test_export_never_overwrites_a_user_file_named_like_a_profile(workdir, name):
    # e.g. a CRM export called customers_profile.csv next to customers.csv
    user_file = workdir / name
    shutil.copy(workdir / name.replace("_profile", ""), user_file)
    before = user_file.read_bytes()
    source = workdir / name.replace("_profile", "")
    result = call("export_profile", file_path=str(source), detect_pii=False)
    assert "won't be overwritten" in error_text(result)
    assert user_file.read_bytes() == before


def test_export_target_that_is_a_folder(workdir):
    (workdir / "customers_profile.csv").mkdir()
    result = call(
        "export_profile", file_path=str(workdir / "customers.csv"), detect_pii=False
    )
    assert "is a folder" in error_text(result)


@posix_only
def test_export_refuses_a_link_to_the_input(workdir):
    source = workdir / "customers.csv"
    (workdir / "customers_profile.csv").symlink_to(source)
    os.link(source, workdir / "hard_profile.csv")
    for output in (None, "hard_profile.csv"):
        result = call(
            "export_profile",
            file_path=str(source),
            output_path=output,
            detect_pii=False,
        )
        assert "would overwrite the input file" in error_text(result)
    assert source.read_bytes() == CUSTOMERS_CSV.read_bytes()


@posix_only
def test_export_refuses_a_link_to_another_user_file(workdir):
    important = workdir / "important.csv"
    important.write_text("keep me")
    (workdir / "customers_profile.csv").symlink_to(important)
    result = call(
        "export_profile", file_path=str(workdir / "customers.csv"), detect_pii=False
    )
    assert "won't be overwritten" in error_text(result)
    assert important.read_text() == "keep me"


@posix_only
def test_linked_input_is_named_and_saved_as_given(workdir):
    (workdir / "reports").mkdir()
    link = workdir / "reports" / "latest.csv"
    link.symlink_to(workdir / "customers.csv")
    data = ok_json(call("profile_dataset", file_path=str(link), detect_pii=False))
    assert data["file_name"] == "latest.csv"
    out = exported(call("export_profile", file_path=str(link), detect_pii=False))
    assert out == workdir / "reports" / "latest_profile.csv"


def test_failed_write_keeps_the_earlier_profile(workdir, monkeypatch):
    source = str(workdir / "customers.csv")
    earlier = exported(call("export_profile", file_path=source, detect_pii=False))
    before = earlier.read_bytes()

    def disk_full(*args, **kwargs):
        raise OSError(28, "No space left on device")

    monkeypatch.setattr(mcp_server.os, "replace", disk_full)
    result = call("export_profile", file_path=source, detect_pii=False)
    assert "No space left on device" in error_text(result)
    assert earlier.read_bytes() == before
    assert not list(workdir.glob(".*.tmp"))


@pytest.mark.parametrize(
    ("output_path", "message"),
    [
        ("profile.xlsx", "must end in .csv"),
        ("customers.csv", "would overwrite the input file"),
        ("missing/profile.csv", "doesn't exist"),
        ("important.csv", "won't be overwritten"),
    ],
)
def test_export_refuses_unsafe_output_paths(workdir, output_path, message):
    (workdir / "important.csv").write_text("keep me")
    result = call(
        "export_profile",
        file_path=str(workdir / "customers.csv"),
        output_path=output_path,
        detect_pii=False,
    )
    assert message in error_text(result)
    assert (workdir / "important.csv").read_text() == "keep me"
    assert (workdir / "customers.csv").read_bytes() == CUSTOMERS_CSV.read_bytes()


@pytest.mark.full_mode
def test_full_export_masks_sensitive_samples(workdir):
    out = exported(call("export_profile", file_path=str(workdir / "company_data.xlsx")))
    rows = list(load_workbook(out)["Customers"].iter_rows(values_only=True))
    header = rows[0]
    by_name = {row[0]: dict(zip(header, row, strict=True)) for row in rows[1:]}
    assert by_name["ssn"]["sample_values"] == "[masked: SSN]"
    assert by_name["email"]["sample_values"] == "[masked: EMAIL_ADDRESS]"
    text = "\n".join(str(v) for row in rows for v in row if v is not None)
    assert_no_raw_sensitive_values(text)


# --- Errors ---------------------------------------------------------------------


@pytest.mark.parametrize("name", ["notes.txt", "old.xls", "data.json", "README"])
@pytest.mark.parametrize("tool", ["profile_dataset", "export_profile"])
def test_only_csv_and_xlsx_paths_are_accepted(tmp_path, tool, name):
    (tmp_path / name).write_text("a,b\n1,2\n")
    message = error_text(call(tool, file_path=str(tmp_path / name)))
    assert "Only .csv and .xlsx files are supported" in message


def test_missing_file(tmp_path):
    message = error_text(call("profile_dataset", file_path=str(tmp_path / "nope.csv")))
    assert "File not found" in message


def test_folder_with_a_csv_name(tmp_path):
    (tmp_path / "folder.csv").mkdir()
    assert "is not a regular file" in error_text(
        call("profile_dataset", file_path=str(tmp_path / "folder.csv"))
    )


def test_permission_errors_are_explained_without_a_traceback(monkeypatch, caplog):
    def denied(self):
        raise PermissionError(13, "Permission denied")

    monkeypatch.setattr(Path, "exists", denied)
    with caplog.at_level(logging.DEBUG):
        message = error_text(call("profile_dataset", file_path=str(CUSTOMERS_CSV)))
    assert "Permission denied" in message
    assert "Privacy & Security" in message
    assert "Traceback" not in caplog.text


def test_unknown_home_folder(tmp_path):
    message = error_text(
        call("profile_dataset", file_path="~no_such_user_dhc/data.csv")
    )
    assert "Use a full path" in message


@pytest.mark.parametrize(
    ("data", "name", "message"),
    [
        (b"", "empty.csv", "empty"),
        (PASSWORD_XLSX.read_bytes(), "locked.xlsx", "password-protected"),
        (b"name\n\x81\x8d\n", "odd.csv", "CSV UTF-8"),
    ],
)
def test_bad_files_give_friendly_errors(tmp_path, data, name, message):
    (tmp_path / name).write_bytes(data)
    text = error_text(
        call("profile_dataset", file_path=str(tmp_path / name), detect_pii=False)
    )
    assert message in text


def test_unexpected_errors_reveal_nothing(monkeypatch, caplog):
    secret = "900-12-3456 alice@example.com"

    def explode(*args, **kwargs):
        raise ValueError(f"could not convert {secret!r}")

    monkeypatch.setattr(profiler, "profile_file", explode)
    with caplog.at_level(logging.DEBUG):
        message = error_text(call("profile_dataset", file_path=str(CUSTOMERS_CSV)))
    assert "Something went wrong" in message
    assert secret not in message
    assert secret not in caplog.text
    assert "ValueError" in caplog.text


def test_files_over_the_size_limit_are_not_read(tmp_path, monkeypatch):
    big = tmp_path / "big.csv"
    big.write_text("a\n1\n")
    monkeypatch.setattr(profiler, "MAX_FILE_BYTES", 2)
    reads = []
    monkeypatch.setattr(Path, "read_bytes", lambda self: reads.append(self) or b"")
    message = error_text(call("profile_dataset", file_path=str(big)))
    assert "over the" in message and "split it" in message
    assert "README" not in message  # web-app advice doesn't apply here
    assert reads == []


# --- Over stdio, as Claude Desktop and Claude Code run it -----------------------


def stdio_server() -> StdioServerParameters:
    # Point the server at the code under test, whatever is installed; -P keeps
    # the working folder off sys.path, as in the documented launch command.
    source_root = str(Path(data_health_check.__file__).parents[1])
    return StdioServerParameters(
        command=sys.executable,
        args=["-P", "-m", "data_health_check.mcp_server"],
        env={**get_default_environment(), "PYTHONPATH": source_root},
    )


def test_stdio_server_ignores_modules_in_the_working_folder(tmp_path):
    # A user's csv.py in the folder the client launches from must not be
    # imported instead of the standard library's csv module.
    (tmp_path / "csv.py").write_text("print('shadow'); raise SystemExit(3)\n")
    params = stdio_server()
    params.cwd = str(tmp_path)
    data = ok_json(
        call(
            "profile_dataset",
            target=params,
            file_path=str(CUSTOMERS_CSV),
            detect_pii=False,
        )
    )
    assert data["row_count"] == 200


def test_stdio_server_lite():
    result = call(
        "profile_dataset",
        target=stdio_server(),
        file_path=str(COMPANY_XLSX),
        detect_pii=False,
        include_samples=True,
    )
    data = ok_json(result)
    assert data["sheet_count"] == 2
    assert_no_raw_sensitive_values(result.content[0].text)


@pytest.mark.full_mode
def test_stdio_server_full():
    # Loads TensorFlow in the server process: its console output must not
    # corrupt the protocol on stdout.
    result = call(
        "profile_dataset",
        target=stdio_server(),
        file_path=str(CUSTOMERS_CSV),
        include_samples=True,
    )
    data = ok_json(result)
    ssn = next(c for c in columns(data) if c["column_name"] == "ssn")
    assert ssn["sample_values"] == "[masked: SSN]"
    assert_no_raw_sensitive_values(result.content[0].text)


def test_lite_mode_hides_min_max_of_number_columns(tmp_path):
    path = tmp_path / "ids.csv"
    path.write_bytes(digits_only_ids_csv())
    result = call("profile_dataset", file_path=str(path), detect_pii=False)
    for column in columns(ok_json(result)):
        assert column["min"] == column["max"] == masking.LITE_MODE_PLACEHOLDER
    assert "900000000" not in result.content[0].text


@pytest.mark.full_mode
def test_full_mode_masks_min_max_of_numeric_ssns(tmp_path):
    path = tmp_path / "ids.csv"
    path.write_bytes(digits_only_ids_csv())
    result = call("profile_dataset", file_path=str(path))
    by_name = {c["column_name"]: c for c in columns(ok_json(result))}
    assert by_name["ssn"]["data_label"] == "SSN"
    assert by_name["ssn"]["min"] == by_name["ssn"]["max"] == "[masked: SSN]"
    assert by_name["quantity"]["max"] == 12
    for raw in ("900000000", "2025550100", "4111111111111111"):
        assert raw not in result.content[0].text


# --- export_profile: the target is checked again right before writing --------


def export_lite(source: Path, **arguments):
    return call("export_profile", file_path=str(source), detect_pii=False, **arguments)


@posix_only
def test_a_fifo_at_the_target_is_refused_without_hanging(workdir):
    os.mkfifo(workdir / "customers_profile.csv")
    assert "won't be overwritten" in error_text(export_lite(workdir / "customers.csv"))


def test_an_annotated_workbook_profile_is_kept(workdir):
    source = workdir / "company_data.xlsx"
    profile_path = exported(export_lite(source))
    workbook = load_workbook(profile_path)
    workbook.create_sheet("My notes", 1)["A1"] = "call the vendor"
    workbook.save(profile_path)
    before = profile_path.read_bytes()
    assert "won't be overwritten" in error_text(export_lite(source))
    assert profile_path.read_bytes() == before


def test_a_profile_resaved_with_semicolons_is_still_recognised(workdir):
    source = workdir / "customers.csv"
    profile_path = exported(export_lite(source))
    text = profile_path.read_text(encoding="utf-8-sig").replace(",", ";")
    profile_path.write_text(text, encoding="utf-8")
    assert exported(export_lite(source)) == profile_path
    assert profile_path.read_text(encoding="utf-8-sig").startswith("column_name,")


def test_a_file_that_appears_while_profiling_is_left_alone(workdir, monkeypatch):
    target = workdir / "customers_profile.csv"
    real_profile = mcp_server._profile

    def slow_profile(*args):
        target.write_text("saved by the user meanwhile")
        return real_profile(*args)

    monkeypatch.setattr(mcp_server, "_profile", slow_profile)
    message = error_text(export_lite(workdir / "customers.csv"))
    assert "while the profile was being made" in message
    assert target.read_text() == "saved by the user meanwhile"


def test_an_earlier_profile_changed_while_profiling_is_kept(workdir, monkeypatch):
    source = workdir / "customers.csv"
    target = exported(export_lite(source))
    real_profile = mcp_server._profile

    def slow_profile(*args):
        with target.open("a", encoding="utf-8") as handle:
            handle.write("my,own,notes\n")
        return real_profile(*args)

    monkeypatch.setattr(mcp_server, "_profile", slow_profile)
    assert "while the profile was being made" in error_text(export_lite(source))
    assert target.read_text(encoding="utf-8-sig").endswith("my,own,notes\n")
    assert not list(workdir.glob(".dhc-*"))


@posix_only
def test_permissions_follow_umask_and_are_kept_on_replace(workdir):
    source = workdir / "customers.csv"
    umask = os.umask(0)
    os.umask(umask)
    target = exported(export_lite(source))
    assert target.stat().st_mode & 0o777 == 0o666 & ~umask
    target.chmod(0o640)
    exported(export_lite(source))
    assert target.stat().st_mode & 0o777 == 0o640


@posix_only
def test_a_read_only_earlier_profile_is_not_replaced(workdir):
    source = workdir / "customers.csv"
    target = exported(export_lite(source))
    target.chmod(0o444)
    assert "read-only" in error_text(export_lite(source))


def test_long_profile_names_can_be_replaced(workdir):
    name = "q" * (250 - len(".csv")) + ".csv"
    source = workdir / "customers.csv"
    first = exported(export_lite(source, output_path=name))
    assert exported(export_lite(source, output_path=name)) == first


@posix_only
def test_link_loops_dangling_links_and_other_targets_are_explained(workdir):
    source = workdir / "customers.csv"
    loop = workdir / "customers_profile.csv"
    loop.symlink_to(loop)
    assert "points back to itself" in error_text(export_lite(source))
    loop.unlink()
    loop.symlink_to(workdir / "elsewhere" / "missing.csv")
    assert "doesn't exist, so nothing was written" in error_text(export_lite(source))
    loop.unlink()
    (workdir / "notes.txt").write_text("keep me")
    loop.symlink_to(workdir / "notes.txt")
    assert "isn't a .csv file" in error_text(export_lite(source))
    assert (workdir / "notes.txt").read_text() == "keep me"


def test_nul_characters_in_paths_are_explained():
    message = error_text(call("profile_dataset", file_path="/tmp/bad\x00name.csv"))
    assert "isn't allowed" in message


def test_headerless_file_does_not_expose_its_first_record(tmp_path):
    path = tmp_path / "raw.csv"
    path.write_bytes(without_header(CUSTOMERS_CSV.read_bytes()))
    result = call("profile_dataset", file_path=str(path), detect_pii=False)
    data = ok_json(result)
    names = [c["column_name"] for c in columns(data)]
    assert names[0] == "Column 1 (no header)"
    assert "no header" in data["sheets"][0]["note"]
    assert_no_raw_sensitive_values(result.content[0].text)
    assert data["warnings"] == []


@pytest.mark.full_mode
def test_unidentified_columns_are_hidden_with_a_warning(tmp_path):
    path = tmp_path / "contacts.csv"
    path.write_bytes(mixed_contacts_csv())
    result = call("profile_dataset", file_path=str(path), include_samples=True)
    data = ok_json(result)
    by_name = {c["column_name"]: c for c in columns(data)}
    assert by_name["contact"]["data_label"] == "could not determine"
    assert by_name["contact"]["sample_values"] == "[hidden: could not determine]"
    assert len(data["warnings"]) == 1
    assert "could not be identified" in data["warnings"][0]
    assert "'contact'" in data["warnings"][0]
    assert_no_raw_sensitive_values(result.content[0].text)


def installed_launcher() -> Path | None:
    folder = Path(sys.executable).parent
    for name in ("data-health-check-mcp", "data-health-check-mcp.exe"):
        if (folder / name).is_file():
            return folder / name
    return None


@pytest.mark.skipif(installed_launcher() is None, reason="package not installed")
def test_installed_command_starts_the_server_safely(tmp_path):
    # The data-health-check-mcp launcher needs no -P: it never puts the folder
    # it is started from on sys.path, so a csv.py there is ignored.
    (tmp_path / "csv.py").write_text("print('shadow'); raise SystemExit(3)\n")
    params = StdioServerParameters(
        command=str(installed_launcher()), args=[], cwd=str(tmp_path)
    )
    data = ok_json(
        call(
            "profile_dataset",
            target=params,
            file_path=str(CUSTOMERS_CSV),
            detect_pii=False,
        )
    )
    assert data["row_count"] == 200
