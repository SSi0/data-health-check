"""Write a profile in the same format as the input file (CSV or XLSX)."""

from __future__ import annotations

import csv
import io
import re
import stat
import warnings
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path, PurePath
from typing import Any
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.drawing.spreadsheet_drawing import SpreadsheetDrawing
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet._writer import WorksheetWriter
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.writer.excel import ExcelWriter

from data_health_check.profiler import (
    COLUMN_FIELDS,
    FIELD_GUIDE,
    LABEL_FIELDS,
    DatasetProfile,
)

CSV_MIME_TYPE = "text/csv"
XLSX_MIME_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

CREATOR = "Data Health Check"
SUMMARY_SHEET = "Summary"
FIELD_GUIDE_SHEET = "Field guide"
FIELD_GUIDE_HEADERS = ("Field", "What it means")
SUMMARY_HEADERS = (
    "File name",
    "Sheet name",
    "Rows",
    "Columns",
    "Rows profiled",
    "Profiled at (UTC)",
    "Mode",
    "Notes",
)

_EXCEL_MAX_CELL_CHARS = 32767
_EXCEL_MAX_SHEET_NAME = 31
_EXCEL_SHEET_NAME_FORBIDDEN = str.maketrans({c: "_" for c in "[]:*?/\\"})
_MIN_WIDTH, _MAX_WIDTH = 8, 60
# Cells starting with these can be run as formulas by spreadsheet apps.
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


@dataclass(frozen=True)
class ExportedProfile:
    file_name: str
    data: bytes
    mime_type: str


def output_file_name(input_name: str) -> str:
    """``customers.xlsx`` -> ``customers_profile.xlsx``."""
    path = PurePath(input_name)
    return f"{path.stem}_profile{path.suffix.lower()}"


def is_profile_file(path: Path) -> bool:
    """True if ``path`` holds a profile written by this tool, unchanged in shape.

    Judged by content, not by name, and used to decide whether a file may be
    replaced: a CSV whose header row is exactly a profile header (also when
    re-saved with ";" or tab separators), or a workbook made by Data Health
    Check whose sheets are Summary, profile sheets and Field guide, each with
    its own header. A workbook someone added a sheet or a column to doesn't
    qualify, so their additions are never thrown away. Never blocks: special
    files (FIFOs, devices) and oversized files are rejected without reading.
    """
    try:
        info = path.stat()
        if not stat.S_ISREG(info.st_mode) or info.st_size > _MAX_PROFILE_BYTES:
            return False
        suffix = path.suffix.lower()
        if suffix == ".csv":
            return _is_profile_csv(path)
        if suffix == ".xlsx":
            return _is_profile_workbook(path)
    except Exception:  # unreadable, not a zip, not text: not one of ours
        return False
    return False


# Profiles are small (one row per input column); anything bigger isn't one.
_MAX_PROFILE_BYTES = 10 * 1024 * 1024
_MAX_UNPACKED_BYTES = 50 * 1024 * 1024
_PROFILE_HEADERS = frozenset(
    {
        COLUMN_FIELDS,
        tuple(name for name in COLUMN_FIELDS if name not in LABEL_FIELDS),
    }
)


def _is_profile_csv(path: Path) -> bool:
    with path.open("rb") as handle:
        header = handle.readline(4096).decode("utf-8-sig").strip()
    return tuple(re.split(r"[,;\t]", header)) in _PROFILE_HEADERS


def _is_profile_workbook(path: Path) -> bool:
    with ZipFile(path) as archive:
        # Cheap checks first: the creator, and no part that would be costly to
        # parse (guards against huge or crafted workbooks at the target).
        if sum(item.file_size for item in archive.infolist()) > _MAX_UNPACKED_BYTES:
            return False
        core = archive.read("docProps/core.xml")
    if f"<dc:creator>{CREATOR}</dc:creator>".encode() not in core:
        return False
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        workbook = load_workbook(path, read_only=True)
    try:
        names = workbook.sheetnames
        if (
            len(names) < 2
            or names[0] != SUMMARY_SHEET
            or names[-1] != FIELD_GUIDE_SHEET
        ):
            return False
        if _first_row(workbook[SUMMARY_SHEET]) != SUMMARY_HEADERS:
            return False
        if _first_row(workbook[FIELD_GUIDE_SHEET]) != FIELD_GUIDE_HEADERS:
            return False
        return all(
            _first_row(workbook[name]) in _PROFILE_HEADERS for name in names[1:-1]
        )
    finally:
        workbook.close()


def _first_row(sheet: Any) -> tuple[str, ...]:
    row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
    return tuple("" if value is None else str(value) for value in row)


def export_profile(
    profile: DatasetProfile, *, reveal_sensitive: bool = False
) -> ExportedProfile:
    """Render ``profile`` in the input's format, masking sensitive samples."""
    name = output_file_name(profile.file_name)
    if profile.file_format == "csv":
        return ExportedProfile(
            name,
            to_csv_bytes(profile, reveal_sensitive=reveal_sensitive),
            CSV_MIME_TYPE,
        )
    return ExportedProfile(
        name, to_xlsx_bytes(profile, reveal_sensitive=reveal_sensitive), XLSX_MIME_TYPE
    )


# --------------------------------------------------------------------------
# CSV
# --------------------------------------------------------------------------


def to_csv_bytes(profile: DatasetProfile, *, reveal_sensitive: bool = False) -> bytes:
    """One row per input column. UTF-8 with a BOM so Excel shows accents."""
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=list(profile.column_fields()))
    writer.writeheader()
    for sheet in profile.sheets:
        for row in profile.column_rows(sheet, reveal_sensitive=reveal_sensitive):
            writer.writerow({key: _csv_safe(value) for key, value in row.items()})
    return buffer.getvalue().encode("utf-8-sig")


def _csv_safe(value: Any) -> Any:
    """Prefix text that a spreadsheet app could run as a formula with ``'``."""
    if (
        isinstance(value, str)
        and value.startswith(_FORMULA_PREFIXES)
        and not _is_numeric_text(value)
    ):
        return "'" + value
    return value


def _is_numeric_text(value: str) -> bool:
    """True for ``-12.5`` or ``-3; 4.25`` (numbers that only look like formulas)."""
    try:
        for part in value.split(";"):
            float(part)
    except ValueError:
        return False
    return True


# --------------------------------------------------------------------------
# XLSX
# --------------------------------------------------------------------------


def to_xlsx_bytes(profile: DatasetProfile, *, reveal_sensitive: bool = False) -> bytes:
    """Summary sheet, one profile sheet per input sheet, and a Field guide."""
    workbook = Workbook()
    workbook.properties.creator = CREATOR

    summary = workbook.active
    summary.title = SUMMARY_SHEET
    _write_table(
        summary,
        SUMMARY_HEADERS,
        [
            (
                profile.file_name,
                sheet.sheet_name,
                sheet.row_count,
                sheet.column_count,
                sheet.rows_profiled,
                profile.profiled_at,
                profile.mode,
                sheet.note,
            )
            for sheet in profile.sheets
        ],
    )

    used_names = {SUMMARY_SHEET.lower(), FIELD_GUIDE_SHEET.lower()}
    fields = profile.column_fields()
    for sheet in profile.sheets:
        title = _unique_sheet_title(sheet.sheet_name or "Profile", used_names)
        rows = profile.column_rows(sheet, reveal_sensitive=reveal_sensitive)
        _write_table(
            workbook.create_sheet(title),
            fields,
            [[row.get(name) for name in fields] for row in rows],
        )

    guide = workbook.create_sheet(FIELD_GUIDE_SHEET)
    _write_table(
        guide,
        FIELD_GUIDE_HEADERS,
        [(name, FIELD_GUIDE[name]) for name in fields],
        wrap_columns=(2,),
        max_width=100,
    )

    return _save_in_memory(workbook)


class _InMemoryExcelWriter(ExcelWriter):
    """openpyxl's ExcelWriter, with worksheets built in memory.

    The stock writer serialises every worksheet to a temporary file on disk
    first. Profiles contain column names and sample values, so they stay in
    memory instead. Mirrors ``ExcelWriter.write_worksheet`` (openpyxl 3.1).
    """

    def write_worksheet(self, ws: Worksheet) -> None:
        ws._drawing = SpreadsheetDrawing()
        ws._drawing.charts = ws._charts
        ws._drawing.images = ws._images
        writer = WorksheetWriter(ws, out=io.BytesIO())
        writer.write()
        ws._rels = writer._rels
        self._archive.writestr(ws.path[1:], writer.read())
        self.manifest.append(ws)


def _save_in_memory(workbook: Workbook) -> bytes:
    buffer = io.BytesIO()
    workbook.properties.modified = datetime.now(UTC).replace(tzinfo=None)
    with ZipFile(buffer, "w", ZIP_DEFLATED, allowZip64=True) as archive:
        _InMemoryExcelWriter(workbook, archive).save()
    return buffer.getvalue()


def _write_table(
    sheet: Worksheet,
    headers: tuple[str, ...] | list[str],
    rows: list[Any],
    *,
    wrap_columns: tuple[int, ...] = (),
    max_width: int = _MAX_WIDTH,
) -> None:
    """Write a table with a bold, frozen header row and fitted column widths."""
    sheet.append(list(headers))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"

    widths = [len(str(header)) for header in headers]
    for row_index, row in enumerate(rows, start=2):
        for col_index, value in enumerate(row, start=1):
            cell = sheet.cell(row=row_index, column=col_index, value=_xlsx_safe(value))
            if isinstance(value, str) and cell.data_type == "f":
                cell.data_type = "s"  # text that starts with "=" stays text
            if col_index in wrap_columns:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            if value is not None:
                widths[col_index - 1] = max(widths[col_index - 1], len(str(value)))

    for col_index, width in enumerate(widths, start=1):
        letter = get_column_letter(col_index)
        sheet.column_dimensions[letter].width = min(
            max(width + 2, _MIN_WIDTH), max_width
        )


def _xlsx_safe(value: Any) -> Any:
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub("", value)[:_EXCEL_MAX_CELL_CHARS]
    return value


def _unique_sheet_title(name: str, used_lower: set[str]) -> str:
    """Make a valid, unique Excel sheet title (31 chars, no []:*?/\\)."""
    base = ILLEGAL_CHARACTERS_RE.sub("", name).translate(_EXCEL_SHEET_NAME_FORBIDDEN)
    base = base.strip().strip("'")[:_EXCEL_MAX_SHEET_NAME] or "Profile"
    title, counter = base, 2
    while title.lower() in used_lower:
        suffix = f" ({counter})"
        title = base[: _EXCEL_MAX_SHEET_NAME - len(suffix)] + suffix
        counter += 1
    used_lower.add(title.lower())
    return title
